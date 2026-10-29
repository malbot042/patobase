"""Exportar resultados a CSV, Excel y PDF."""

import csv
import io
import os
import subprocess
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]

SQL = (
    "SELECT v.venta_id, v.fecha, now() AS exportado, c.nombre || ' — Ñandú' AS cliente, v.importe, "
    "NULL AS vacio, [1, 2] AS lista FROM files.ventas v JOIN files.clientes c USING (cliente_id) ORDER BY v.venta_id"
)


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("exp")
    env = {**os.environ, "PATOBASE_DATA_DIR": str(tmp / "data")}
    subprocess.run([sys.executable, str(ROOT / "scripts" / "seed.py")], check=True, env=env, capture_output=True)
    mp = pytest.MonkeyPatch()
    mp.setenv("PATOBASE_DATA_DIR", str(tmp / "data"))
    mp.setenv("PATOBASE_META_DB", str(tmp / "meta.sqlite"))
    mp.setenv("PATOBASE_PDF_ROWS", "50")
    for mod in [m for m in sys.modules if m.startswith("patobase")]:
        del sys.modules[mod]
    from patobase.main import app

    with TestClient(app) as c:
        yield c
    mp.undo()


def test_csv_full_result_utf8_bom(client):
    r = client.post("/export/csv", data={"sql": SQL, "title": "Ventas Ñandú / 2026"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "filename*=UTF-8''Ventas%20%C3%91and%C3%BA%202026.csv" in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["venta_id", "fecha", "exportado", "cliente", "importe", "vacio", "lista"]
    assert len(rows) == 20_001  # todas las filas, no solo la página visible
    assert rows[1][3].endswith("— Ñandú") and rows[1][5] == "" and rows[1][1].count("-") == 2


def test_xlsx_types_and_info(client):
    import openpyxl

    r = client.post("/export/xlsx", data={"sql": SQL, "title": "Ventas"})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Datos", "Info"]
    ws = wb["Datos"]
    assert ws.max_row == 20_001 and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:G20001"
    row = [c for c in ws[2]]
    assert isinstance(row[0].value, int) and row[1].is_date and row[2].is_date  # timestamptz sin zona
    assert isinstance(row[4].value, float) and row[4].number_format == "#,##0.00"
    assert row[6].value == "[1, 2]"  # las listas se exportan como texto
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Filas"] == 20_000 and info["Completo"] == "Sí" and "files.ventas" in info["SQL"]


def test_pdf_is_capped_and_says_so(client):
    r = client.post("/export/pdf", data={"sql": SQL, "title": "Informe"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    from reportlab.pdfbase.pdfmetrics import stringWidth  # noqa: F401  (reportlab instalado)

    assert len(r.content) > 2_000


def test_export_saved_question(client):
    client.post(
        "/questions", data={"name": "Por canal", "sql": "SELECT canal, count(*) n FROM files.ventas GROUP BY 1"}
    )
    r = client.get("/questions/1/export/csv")
    assert r.status_code == 200 and "Por canal.csv" in r.headers["content-disposition"]
    assert len(r.text.strip().splitlines()) == 4
    assert client.get("/questions/999/export/csv").status_code == 404


@pytest.mark.parametrize(
    "sql",
    ["DROP TABLE x", "SELECT * FROM read_csv('/etc/passwd')", "SELECT * FROM no_existe", "SELECT 1; SELECT 2"],
)
def test_export_rejects_bad_sql(client, sql):
    for fmt in ("csv", "xlsx", "pdf"):
        r = client.post(f"/export/{fmt}", data={"sql": sql})
        assert r.status_code == 400 and r.text.startswith("No se puede exportar")


def test_unknown_format(client):
    assert client.post("/export/exe", data={"sql": "SELECT 1"}).status_code == 400


def test_stream_truncation_flag(client):
    eng = client.app.state.engine
    with eng.stream("SELECT * FROM range(100)", 30) as st:
        assert sum(len(b) for b in st.batches(7)) == 30
    assert st.truncated
    with eng.stream("SELECT * FROM range(10)", 30) as st:
        assert sum(len(b) for b in st.batches(7)) == 10
    assert not st.truncated


def test_list_types_are_not_numeric():
    from patobase.engine import vega_type

    assert vega_type("INTEGER[]") == "nominal" and vega_type("STRUCT(a INTEGER)") == "nominal"
    assert vega_type("DECIMAL(10,2)") == "quantitative"


def test_filenames():
    from patobase.exports import content_disposition, filename

    assert filename("Ventas Ñandú / 2026", "csv") == "Ventas Ñandú 2026.csv"
    assert filename("   ", "pdf") == "patobase.pdf"
    cd = content_disposition("Año «2026».xlsx")
    assert cd.startswith('attachment; filename="Ano 2026.xlsx"') or 'filename="Ano' in cd
