"""Exportar resultados a CSV, Excel (.xlsx) y PDF.

- CSV: UTF-8 con BOM (Excel lo abre con acentos correctos), separador coma,
  decimales con punto: el formato estándar que lee cualquier herramienta.
  Se genera en streaming, sin cargar el resultado en memoria.
- Excel: celdas con su tipo real (números, fechas), cabecera fija, autofiltro
  y una hoja «Info» con la consulta, la fecha y si el resultado se recortó.
- PDF: informe con título, fecha y tabla paginada (cabecera repetida en cada
  página), números en formato español. Pensado para leer, no para reprocesar:
  se limita a PATOBASE_PDF_ROWS filas.
"""

from __future__ import annotations

import csv
import io
import itertools
import re
import unicodedata
from collections.abc import Iterator
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from typing import Any
from urllib.parse import quote

from .engine import RowStream, vega_type

FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}
EXCEL_MAX_ROWS = 1_048_575  # 1 048 576 filas por hoja, menos la cabecera


def filename(title: str, ext: str) -> str:
    base = unicodedata.normalize("NFKC", title or "").strip() or "patobase"
    base = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", base)
    base = re.sub(r"\s+", " ", base).strip()[:80] or "patobase"
    return f"{base}.{ext}"


def content_disposition(name: str) -> str:
    """Nombre de fichero que funciona en todos los navegadores, también con acentos."""
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode() or "export"
    ascii_name = ascii_name.replace('"', "")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}"


def text(v: Any) -> str:
    """Valor como texto neutro (ISO para fechas, punto decimal)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, bytes):
        return v.hex()
    return str(v)


# --------------------------------------------------------------------- CSV
def csv_chunks(stream: RowStream, batch: int = 5_000) -> Iterator[bytes]:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    yield "﻿".encode()  # BOM: Excel detecta UTF-8
    w.writerow(stream.columns)
    for rows in stream.batches(batch):
        for r in rows:
            w.writerow([text(v) for v in r])
        yield buf.getvalue().encode("utf-8")
        buf.seek(0)
        buf.truncate()
    if buf.tell():
        yield buf.getvalue().encode("utf-8")


# ------------------------------------------------------------------- Excel
def _excel_value(v: Any) -> Any:
    if v is None or isinstance(v, (bool, int, float, str)):
        if isinstance(v, int) and not isinstance(v, bool) and abs(v) > 2**53:
            return str(v)  # Excel perdería precisión
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.astimezone().replace(tzinfo=None) if v.tzinfo else v  # Excel no admite zona horaria
    if isinstance(v, time):
        return v.replace(tzinfo=None)
    if isinstance(v, (date, timedelta)):
        return v
    return text(v)


def to_xlsx(stream: RowStream, title: str, sql: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Datos")
    ws.freeze_panes = "A2"
    vtypes = [vega_type(t) for t in stream.types]
    widths = [min(max(len(c), 8), 60) for c in stream.columns]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="D4890B")

    header = []
    for c in stream.columns:
        cell = WriteOnlyCell(ws, value=c)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center")
        header.append(cell)

    # Formatos por tipo de columna
    fmts = []
    for t, vt in zip(stream.types, vtypes, strict=True):
        tu = t.upper()
        if tu.startswith("DATE"):
            fmts.append("dd/mm/yyyy")
        elif tu.startswith("TIMESTAMP"):
            fmts.append("dd/mm/yyyy hh:mm:ss")
        elif tu.startswith(("DECIMAL", "DOUBLE", "FLOAT")):
            fmts.append("#,##0.00")
        elif vt == "quantitative":
            fmts.append("#,##0")
        else:
            fmts.append(None)

    # Las anchuras deben fijarse antes de escribir filas en modo write_only:
    # se estiman con la cabecera y la primera tanda de datos.
    batches = stream.batches(5_000)
    first = next(batches, [])
    for r in first[:500]:
        for i, v in enumerate(r):
            widths[i] = min(max(widths[i], len(text(v)) + 2), 60)
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.append(header)

    written = 0
    capped = False
    for rows in itertools.chain([first], batches) if first else ():
        for r in rows:
            if written >= EXCEL_MAX_ROWS:
                capped = True
                break
            ws.append([_cell(ws, _excel_value(v), f, WriteOnlyCell) for v, f in zip(r, fmts, strict=True)])
            written += 1
        if capped:
            break
    if stream.columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(stream.columns))}{written + 1}"

    info = wb.create_sheet("Info")
    for row in [
        ["Pregunta", title or "(sin título)"],
        ["Exportado", datetime.now().strftime("%d/%m/%Y %H:%M")],
        ["Filas", written],
        ["Completo", "No: resultado recortado" if (stream.truncated or capped) else "Sí"],
        ["SQL", sql],
    ]:
        info.append(row)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(ws, value, fmt, cell_cls):
    if fmt is None or value is None or isinstance(value, str):
        return value
    c = cell_cls(ws, value=value)
    c.number_format = fmt
    return c


# --------------------------------------------------------------------- PDF
def es_number(v: Any) -> str:
    if isinstance(v, bool) or not isinstance(v, (int, float, Decimal)):
        return text(v)
    s = f"{v:,.2f}" if isinstance(v, (float, Decimal)) else f"{v:,}"
    return s.replace(",", " ").replace(".", ",")


def _pdf_text(v: Any, vt: str) -> str:
    if v is None:
        return ""
    if vt == "quantitative":
        return es_number(v)
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M") if (v.hour or v.minute or v.second) else v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return text(v)


def to_pdf(stream: RowStream, title: str, max_rows: int) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    ncols = max(len(stream.columns), 1)
    pagesize = landscape(A4) if ncols > 5 else A4
    font_size = 8 if ncols <= 8 else 7 if ncols <= 12 else 6
    vtypes = [vega_type(t) for t in stream.types]

    # El stream se abre con max_rows = filas del PDF: al leerlo entero sabe si había más
    rows: list[tuple] = [r for batch in stream.batches(2_000) for r in batch][:max_rows]
    more = stream.truncated

    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=font_size, leading=font_size + 2)
    cell_num = ParagraphStyle("num", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    head_num = ParagraphStyle("headnum", parent=head, alignment=2)

    def para(s: str, style) -> Paragraph:
        s = s if len(s) <= 300 else s[:297] + "…"
        return Paragraph(escape(s), style)

    data = [[para(c, head_num if vt == "quantitative" else head) for c, vt in zip(stream.columns, vtypes, strict=True)]]
    for r in rows:
        data.append(
            [
                para(_pdf_text(v, vt), cell_num if vt == "quantitative" else cell)
                for v, vt in zip(r, vtypes, strict=True)
            ]
        )

    # Anchura de cada columna proporcional a su contenido (cabecera y primeras filas),
    # con un mínimo para que nada quede ilegible y un máximo para que nada lo acapare
    doc_w = pagesize[0] - 24 * mm
    sample = rows[:200]
    lens = [
        min(max([len(c)] + [len(_pdf_text(r[i], vt)) for r in sample]), 40)
        for i, (c, vt) in enumerate(zip(stream.columns, vtypes, strict=True))
    ] or [1]
    lens = [max(n, 4) for n in lens]
    widths = [doc_w * n / sum(lens) for n in lens]
    table = LongTable(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D4890B")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F3EC")]),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2DCCF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    note = f"{len(rows)} fila{'s' if len(rows) != 1 else ''}"
    if more:
        note += f" (primeras {max_rows}; exporta a CSV o Excel para el resultado completo)"
    story = [
        Paragraph(escape(title or "Resultado"), styles["Title"]),
        Paragraph(f"Exportado el {now} · {note}", styles["Normal"]),
        Spacer(1, 6 * mm),
        table,
    ]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 8 * mm, f"patobase · {escape(title or '')}"[:120])
        canvas.drawRightString(pagesize[0] - 12 * mm, 8 * mm, f"Página {doc.page}")
        canvas.restoreState()

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=title or "patobase",
        author="patobase",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
